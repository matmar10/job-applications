#!/usr/bin/env python3
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_page_x_of_y(paragraph):
    """Insert PAGE and NUMPAGES fields for 'Page X of Y'."""
    def make_field(instrText):
        run = OxmlElement('w:r')
        fldChar_begin = OxmlElement('w:fldChar')
        fldChar_begin.set(qn('w:fldCharType'), 'begin')
        run.append(fldChar_begin)
        r2 = OxmlElement('w:r')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = instrText
        r2.append(instr)
        r3 = OxmlElement('w:r')
        fldChar_end = OxmlElement('w:fldChar')
        fldChar_end.set(qn('w:fldCharType'), 'end')
        r3.append(fldChar_end)
        return run, r2, r3

    run = paragraph.add_run('Resume of Matthew Joseph Martin – Staff Software Engineer    Page ')
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    for el in make_field(' PAGE '):
        paragraph._p.append(el)

    mid = paragraph.add_run(' of ')
    mid.font.size = Pt(7)
    mid.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    for el in make_field(' NUMPAGES '):
        paragraph._p.append(el)

def set_footer(docx_path):
    doc = Document(docx_path)
    for section in doc.sections:
        footer = section.footer
        # Clear existing footer paragraphs
        for p in footer.paragraphs:
            p.clear()
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_page_x_of_y(para)
    doc.save(docx_path)
    print(f"Footer added: {docx_path}")

if __name__ == '__main__':
    if len(sys.argv) != 2:
        print("Usage: add-footer.py <path-to.docx>")
        sys.exit(1)
    set_footer(sys.argv[1])
